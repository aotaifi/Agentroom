"""Build a five-slide pitch from the prototype's acceptance report."""
import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(17, 28, 46)
WHITE = RGBColor(238, 244, 249)
MUTED = RGBColor(166, 183, 201)
TEAL = RGBColor(74, 219, 181)


def text(slide, x, y, w, h, value, size=24, color=WHITE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(16)
    return box


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--report", type=Path, required=True)
    p.add_argument("--output", type=Path, required=True)
    a = p.parse_args()
    report = json.loads(a.report.read_text())
    for key in ("isolation", "science", "submission_tests", "generic_isolation"):
        if not report[key]["passed"]:
            raise ValueError("passing live checks are required before claiming readiness")
    r = report["settings"]["resources"]
    slides = [
        ("Agent freedom.\nControl outside.", "A real development workspace for every agent.\nYour host, credentials and cluster controls stay outside.", "AGENTROOM / WORKING PROTOTYPE"),
        ("One pod. One task.", "Profile  →  rootless pod  →  structured job gateway\nCopied project. Pinned dependencies. Familiar shell.\nDestination workers are the next integration step.", "THE PRODUCT"),
        ("Chatty is the first example.", f"{r['cpus']} logical CPUs · {r['memory']} RAM · {r['workspace']} workspace · {r['pids']} processes\nOffline · non-root · read-only image · no keys or host sockets\nPython 3.11 with JAX / NetKet / Flax / SciPy / Numba", "LIVE SETTINGS / WS1"),
        ("Generic by design.", "Resources → CPU, RAM, filesystem capacity, processes\nEnvironment → image, dependencies, explicit variables\nPermissions → copied inputs; enforced isolation\nBackends → independent resource ceilings and worker contracts", "ONE JSON PROFILE / MANY PROJECTS"),
        ("Review now. Integrate next.", "Live pod and isolation checks: passed\nSlurm + workstation-pool preparation: verified\nLive submission: disabled until isolated workers exist\nChatty research state: untouched", "STARTUP PITCH / VERIFIED SCOPE"),
    ]
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for index, (title, body, label) in enumerate(slides, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = NAVY
        text(slide, 0.8, 0.55, 11.7, 0.4, label, 13, TEAL)
        text(slide, 0.8, 1.35, 11.7, 2.0, title, 42)
        text(slide, 0.8, 3.55, 11.7, 2.65, body, 25, MUTED)
        text(slide, 0.8, 6.8, 11.7, 0.3, f"Agentroom · prototype v0.1                                      {index:02} / 05", 12, MUTED)
    a.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(a.output)
    print(a.output)


if __name__ == "__main__":
    main()
